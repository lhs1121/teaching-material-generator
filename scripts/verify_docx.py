# -*- coding: utf-8 -*-
"""verify_docx.py —— 校验生成的 docx 是否满足编写铁律第 8 条（全文宋体）与结构完整。

用法：
    python verify_docx.py "教案 V1.docx" "课件大纲 V1.docx"
    python verify_docx.py --code-font Consolas dist/*.docx   # 转换时用了 --code-font，校验须同步

说明：md2docx.py 支持把代码块（板书 ASCII 树形图）换成等宽字体。若转换时用了
`--code-font`，本脚本须带同名参数，否则代码块会被误判为"非宋体"。
"""
import argparse
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

FONT = "宋体"
# 与 md2docx.py 保持一致：这几个样式属代码块（板书树形图）
CODE_STYLES = ("Source Code", "Preformatted Text")


def all_runs(paragraph):
    """返回段落内所有 run，含超链接中的 run（paragraph.runs 会漏掉后者）。"""
    return [Run(el, paragraph) for el in paragraph._element.iter(qn("w:r"))]


def check(path: Path, code_font: str = FONT):
    doc = Document(str(path))
    print(f"\n=== {path.name} ===")

    paras = doc.paragraphs
    headings = [p for p in paras if p.style is not None and p.style.name.startswith("Heading")]
    print(f"段落 {len(paras)} 个｜标题 {len(headings)} 个｜表格 {len(doc.tables)} 个")

    # 1) 全文字体检查（正文 + 标题 + 表格内）
    #    代码块按 code_font 判定，其余按宋体判定
    bad = []
    total_runs = 0
    code_runs = 0
    for p in paras:
        style_name = p.style.name if p.style is not None else ""
        expected = code_font if style_name in CODE_STYLES else FONT
        for r in all_runs(p):
            if not r.text.strip():
                continue
            total_runs += 1
            if style_name in CODE_STYLES:
                code_runs += 1
            if r.font.name != expected:
                bad.append((p.text[:18], r.font.name, expected))
    cell_runs = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in all_runs(p):
                        if not r.text.strip():
                            continue
                        cell_runs += 1
                        if r.font.name != FONT:
                            bad.append((p.text[:18], r.font.name, FONT))

    print(f"正文/标题 run {total_runs} 个（含代码块 {code_runs} 个）｜表格内 run {cell_runs} 个")
    print(f"字体不符 run：{len(bad)} 个" + (" ✅" if not bad else " ❌"))
    for b in bad[:5]:
        print(f"   - 「{b[0]}」实际={b[1]}｜期望={b[2]}")

    # 2) 表格边框与结构
    print("表格明细：")
    for i, t in enumerate(doc.tables, 1):
        has_border = t._tbl.tblPr.find(qn("w:tblBorders")) is not None
        print(f"  表{i}: {len(t.rows)}行 × {len(t.columns)}列｜边框={'有 ✅' if has_border else '无 ❌'}")

    # 3) 标题层级
    print("标题层级：")
    for p in headings:
        print(f"  [{p.style.name}] {p.text[:36]}")


def main():
    ap = argparse.ArgumentParser(description="校验 docx 字体与结构（编写铁律第 8 条）")
    ap.add_argument("files", nargs="+", help="待校验的 .docx 文件")
    ap.add_argument(
        "--code-font",
        default=FONT,
        help=f"代码块（板书树形图）字体，须与 md2docx.py 的 --code-font 一致，默认 {FONT}",
    )
    args = ap.parse_args()

    for f in args.files:
        p = Path(f)
        if not p.exists():
            print(f"[跳过] 不存在：{p}")
            continue
        check(p, code_font=args.code_font)

    return 0


if __name__ == "__main__":
    sys.exit(main())
